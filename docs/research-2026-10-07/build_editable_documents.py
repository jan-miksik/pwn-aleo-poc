"""Build plain editable DOCX documents from the sibling Markdown sources."""
from pathlib import Path
import re
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
DOCS = ROOT.parent

def inline(paragraph, text, size=None):
    pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|`[^`]+`)'
    for token in re.split(pattern, text):
        if not token:
            continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
        if m:
            link=OxmlElement('w:hyperlink')
            link.set(qn('r:id'),paragraph.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');props=OxmlElement('w:rPr')
            color=OxmlElement('w:color');color.set(qn('w:val'),'24486B');props.append(color)
            underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline)
            if size:
                font_size=OxmlElement('w:sz');font_size.set(qn('w:val'),str(int(size*2)));props.append(font_size)
            run.append(props);t=OxmlElement('w:t');t.text=m[1];run.append(t);link.append(run)
            paragraph._p.append(link)
        else:
            code=token.startswith('`')
            bold=token.startswith('**')
            r=paragraph.add_run(token[1:-1] if code else token[2:-2] if bold else token)
            r.bold=bold
            if code: r.font.name='Courier New';r.font.size=Pt((size or 10.5)-1)
            elif size: r.font.size=Pt(size)

def setup(title):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1.9);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
    s.header_distance=Cm(0.8);s.footer_distance=Cm(0.8)
    styles=d.styles
    normal=styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
    normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',22),('Heading 1',14),('Heading 2',11.5),('Heading 3',11)]:
        st=styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.font.bold=True;st.paragraph_format.space_before=Pt(12 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
    for n in ['List Bullet','List Number']:
        styles[n].font.name='Arial';styles[n].font.size=Pt(10.5)
        styles[n].paragraph_format.space_after=Pt(4)
    h=s.header.paragraphs[0];h.add_run(title).font.size=Pt(8)
    h.paragraph_format.space_after=Pt(0)
    footer=s.footer.paragraphs[0];footer.alignment=2
    r=footer.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
    d.core_properties.title=title;d.core_properties.author='PWN DAO Foundation'
    return d

def table(d, rows):
    n=len(rows[0]);t=d.add_table(rows=1, cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[17/n]*n
    if n==3:
        widths=[4.2,5.4,7.4]
        if rows[0][0] in ['Program','Pořadí']: widths=[9.3,3.5,4.2]
        if rows[0][0]=='Milestone': widths=[11,3,3]
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,value in enumerate(row):
            cell=cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');mar.append(x)
            cp.append(mar)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EDF2');cp.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.07
            inline(p,value,9)
            if i==0:
                for r in p.runs:r.bold=True
        rowprops=t.rows[i]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');rowprops.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');rowprops.append(repeat)
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    d.add_paragraph().paragraph_format.space_after=Pt(0)

def build(path):
    lines=path.read_text().splitlines();title=lines[0][2:]
    d=setup(title);i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in cells):rows.append(cells)
                i+=1
            table(d,rows);continue
        if line.startswith('# '):p=d.add_paragraph(style='Title');inline(p,line[2:])
        elif line.startswith('## '):p=d.add_paragraph(style='Heading 1');inline(p,line[3:])
        elif line.startswith('### '):p=d.add_paragraph(style='Heading 2');inline(p,line[4:])
        elif line.startswith('- '):p=d.add_paragraph(style='List Bullet');inline(p,line[2:])
        else:
            text=line
            while i+1<len(lines) and lines[i+1].strip() and not lines[i+1].startswith(('#','|','- ')):
                i+=1;text+=' '+lines[i].strip()
            p=d.add_paragraph();inline(p,text)
        i+=1
    out=path.with_suffix('.docx');d.save(out);print(out)

def ranking_appendix():
    data=json.loads((ROOT/'ranking.json').read_text())['rankings']
    evidence={x['program']:x for x in json.loads((ROOT/'program-evidence.json').read_text())['programs']}
    out=['# Aleo top programů a aplikační upgrady','', 'Ověřeno 7. října 2026. Zdroj žebříčku: [Provable Explorer](https://beta.explorer.provable.com/programs). Ano znamená cestu upgradu v constructoru; nezaručuje průchod konkrétního upgrade návrhu. Týdenní dataset neobsahuje credits.aleo.','']
    for key,title in [('allTimeData','Top 20 podle celkových calls'),('weeklyData','Top 20 podle týdenních calls')]:
        out.extend(['## '+title,'','| Program | Calls | Upgrade cesta |','| --- | --- | --- |'])
        for row in data[key][:20]:
            item=evidence[row['program_id']];c=item.get('constructor')
            upgrade=bool(c and c!='assert.eq edition 0u16;')
            out.append(f"| {row['program_id']} | {row['calls']:,} | {'Ano' if upgrade else 'Ne'} |")
        out.append('')
    (ROOT/'top-programs-upgrades.md').write_text('\n'.join(out))

if __name__=='__main__':
    ranking_appendix()
    build(DOCS/'private-lending-aleo-grant-draft-v4.md')
    build(DOCS/'aleo-privacy-upgrades-collateral-review-2026-10-07.md')
