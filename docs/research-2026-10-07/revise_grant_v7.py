from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor

D=Path(__file__).resolve().parent.parent
source=(D/'private-lending-aleo-grant-overview-v6.html').read_text()
old='The team previously developed PWN, a peer-to-peer lending protocol deployed on 13 EVM chains, and now focuses on <a href="https://ownlabs.co/">OWN</a>, building on PWN’s original ambition to make crypto-backed mortgages possible. PWN’s contracts remain publicly available lending infrastructure.'
new='The team previously developed PWN, a peer-to-peer lending protocol deployed on 13 EVM chains. The team now operates as <a href="https://ownlabs.co/">OWN</a>, building crypto-backed mortgages on top of PWN’s lending infrastructure.'
assert source.count(old)==1
(D/'private-lending-aleo-grant-overview-v7.html').write_text(source.replace(old,new))
for name in ['render_illustrated_grant','build_illustrated_grant']:
    ext='cjs' if name.startswith('render') else 'py'
    # Reuse the same six unchanged illustrations and fonts.
    script=(D/'research-2026-10-07'/f'{name}_v6.{ext}').read_text().replace('overview-v6','overview-v7').replace('draft-v6','draft-v7')
    (D/'research-2026-10-07'/f'{name}_v7.{ext}').write_text(script)

sections=[
('Úvod a tým',[
('Upraveno','Projekt je představen jako univerzální lending infrastruktura od týmu za OWN, vycházející z PWN v1.5. OWN staví crypto-backed hypotéky na PWN infrastruktuře.'),
('Upraveno','„13 EVM mainnets“ → „13 EVM chains“. Úvodní karta žadatele nyní zmiňuje tým za OWN a návaznost na PWN.'),
('Odebráno','Datum „September 2026“ v úvodu i v privacy sekci.'),
('Přidáno','Tým na začátku dokumentu: Josef Jelacic — Founder of PWN DAO and OWN; Šimon Kozák — Engineer, Founding Engineer & CTO at PWN DAO and OWN; Jan Mikšík — Engineer, 4 years of building PWN’s DeFi lending application. LinkedIn u všech tří, GitHub u Šimona a Jana.')]),
('Model půjček a nabídek',[
('Upraveno','Popis defaultu výslovně uvádí možnost claimu při nesplacení v době splatnosti. Formulace o likvidacích je zpřesněna na „price-based liquidation“ a přínos pro privacy na minimalizaci zveřejňovaných údajů.'),
('Přidáno','Rozdíl proposal vs. loan: nabídka poskytuje financování, každý draw vytváří samostatnou půjčku. Expirace či zrušení nabídky neukončuje existující půjčky.'),
('Upraveno','Více draws už není prezentováno jako větší anonymita. Adresa lendera zůstává skrytá v lending voláních, vazba commitment–loan je veřejná. U direct offers přibyla privátní autorizace určeného borrowera.')]),
('Architektura a rozšiřitelnost',[
('Přidáno','Označení architektury jako cílového produkčního návrhu. Elastic Proposal je první typ nabídky; další kompatibilní proposal a product moduly lze přidávat bez změny nasazeného core.'),
('Přidáno','Borrower a lender hooks z PWN v1.5 jako reference pro budoucí Aleo integrace. Hooks jsou také mezi možnými rozšířeními.')]),
('Audit a upgradeability',[
('Nahrazeno','Dočasná upgradeability a otázka tříměsíčního až šestiměsíčního přechodného období → dokončený externí audit, vyřešení závažných nálezů a následné nasazení core bez upgradeability.'),
('Přidáno','Výslovné rozlišení mezi immutable core a samostatnými issuer či upgrade oprávněními externích tokenových programů.')]),
('Privacy',[
('Odebráno','„Specific terms of an individual loan“ ze seznamu privátních údajů.'),
('Upřesněno','Privátní historie splácení znamená její přiřazení ke konkrétní identitě. Aktivita půjček pod veřejnými identifikátory zůstává viditelná. V tabulce už není tvrzení, že více draws oslabuje vazbu commitment–loan.'),
('Přidáno','„Direct offers can keep more terms private.“')]),
('Milestones a mainnet',[
('Odebráno','Všechny launch caps: z úvodu, Current status, Config, M1, M4, časové osy, samostatného calloutu i Out of scope.'),
('Doplněno','M1 zahrnuje rozhraní pro proposal a product rozšíření a vymezení privacy. M4 vyžaduje hotový audit a vyřešené závažné nálezy před mainnet vklady; zveřejní se také pravidla immutable nasazení.')]),
('Collateral policy a údržba',[
('Upraveno','Původní otázka, zda zavést whitelist, nahrazena otázkou přijatelných issuer oprávnění a řešení deficitu escrow. Text zohledňuje existující Hub allowlist v multi-token PoC i freeze a pause rizika.'),
('Upraveno','Po grantu udržuje tým za OWN SDK, aplikaci a integrace v rámci infrastruktury PWN DAO Foundation; nasazený core zůstává immutable.')]),
('Reference a formát',[
('Přidáno','OWN jako reference týmu a produktu; reference PWN v1.5 nyní míří na větev v1.5-withdraw-restriction a zmiňuje modulární produkty, proposals a hooks. Ostatní uvedené URL jsou klikatelné.'),
('Přidáno','DOCX s editovatelným textem a tabulkami a šesti ilustracemi. Týmová sekce posunula stránkování; konec časové osy má popisek „Mainnet“.')])]

doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(1.7);s.left_margin=s.right_margin=Cm(1.8)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',21),('Heading 1',12)]:
 st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('17191C')
 st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
title='Přehled změn grantového návrhu'
doc.core_properties.title=title
doc.add_paragraph(title,style='Title')
doc.add_paragraph('Změny oproti původnímu draftu před dnešními úpravami (v3 → v7).')
md=['# '+title,'','Změny oproti původnímu draftu před dnešními úpravami (v3 → v7).','']
for heading,items in sections:
 p=doc.add_paragraph(heading,style='Heading 1')
 if heading=='Privacy':p.paragraph_format.page_break_before=True
 md+=['## '+heading,'']
 for label,body in items:
  p=doc.add_paragraph();p.paragraph_format.widow_control=True
  r=p.add_run(label+'. ');r.bold=True;r.font.color.rgb=RGBColor.from_string('2E6E7E')
  p.add_run(body)
  md+=['- **'+label+'.** '+body]
 md.append('')
doc.save(D/'private-lending-aleo-grant-changes-v3-to-v7.docx')
(D/'private-lending-aleo-grant-changes-v3-to-v7.md').write_text('\n'.join(md))
print('Updated v7 HTML and concise change summary')
