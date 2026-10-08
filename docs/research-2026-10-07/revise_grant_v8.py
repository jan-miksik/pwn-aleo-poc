from pathlib import Path
from difflib import SequenceMatcher
from collections import OrderedDict
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from html import escape

D=Path(__file__).resolve().parent.parent
# Reuse the established semantic extractor without running the earlier authoring code.
scope={}
extractor=(D/'research-2026-10-07/build_grant_v6_changes.py').read_text().split('\nbefore=blocks(')[0]
scope['__file__']=str(D/'research-2026-10-07/build_grant_v6_changes.py')
exec(compile(extractor,'block_extractor','exec'),scope)
blocks=scope['blocks']
before=blocks(D.parent/'archive/docs/private-lending-aleo-grant-overview-v3.html')
after=blocks(D/'private-lending-aleo-grant-overview-v8.html')
def by_section(items):
 result=OrderedDict()
 for section,value in items:
  if value==section:continue
  result.setdefault(section,[]).append(value)
 return result
old=by_section(before);new=by_section(after)
order=list(new)+[x for x in old if x not in new]
doc=Document();page=doc.sections[0]
page.page_width=Cm(21);page.page_height=Cm(29.7)
page.top_margin=page.bottom_margin=Cm(1.7);page.left_margin=page.right_margin=Cm(1.8)
st=doc.styles['Normal'];st.font.name='Arial';st.font.size=Pt(10)
st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(7)
for name,size in [('Title',21),('Heading 1',12)]:
 st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('17191C')
 st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
doc.core_properties.title='Změny grantového návrhu'
doc.add_paragraph('Změny grantového návrhu',style='Title')
doc.add_paragraph('Původní draft v3 před úpravami ze 7. října → aktuální v8. Zobrazeny jsou jen změny s krátkým kontextem. Výpustka (…) označuje nezměněný text.')
def run(p,value,kind):
 r=p.add_run(value)
 if kind in ['delete','insert']:
  r.font.color.rgb=RGBColor.from_string('9C2525' if kind=='delete' else '215E37')
  if kind=='delete':r.font.strike=True
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'FADBD8' if kind=='delete' else 'DDF0DF');sh.set(qn('w:val'),'clear');r._r.get_or_add_rPr().append(sh)
 return r
p=doc.add_paragraph();run(p,'Odebrané znění','delete');p.add_run('   ');run(p,'Nové znění','insert')
htmlparts=['<!doctype html><meta charset="utf-8"><title>Změny grantového návrhu</title><style>body{font:15px/1.55 Arial;max-width:900px;margin:40px auto;color:#17191c}h1{font-size:28px}h2{font-size:19px;margin-top:28px}p{white-space:pre-wrap}del{background:#fadbd8;color:#9c2525}ins{background:#ddf0df;color:#215e37;text-decoration:none}</style><h1>Změny grantového návrhu</h1><p>Původní v3 → aktuální v8. Výpustka (…) označuje nezměněný text.</p>']
ledger=[]
def aligned(a,b):
 # Align whole content blocks first. Never match words across paragraphs.
 for op,i,j,k,l in SequenceMatcher(None,a,b,autojunk=False).get_opcodes():
  if op=='equal':continue
  aa=a[i:j];bb=b[k:l]
  if op!='replace':
   for x in aa:yield x,''
   for x in bb:yield '',x
   continue
  # Ordered best matching prevents inserted headings or paragraphs from being
  # paired with an unrelated removed paragraph.
  pairs=[]
  for x in range(len(aa)):
   for y in range(len(bb)):
    ratio=SequenceMatcher(None,aa[x],bb[y],autojunk=False).ratio()
    if aa[x].split()[:3]==bb[y].split()[:3]:ratio+=0.6
    if ratio>=0.30:pairs.append((ratio,x,y))
  selected=[]
  for _,x,y in sorted(pairs,reverse=True):
   if all((x<u and y<v) or (x>u and y>v) for u,v in selected):selected.append((x,y))
  px=py=0
  for x,y in sorted(selected):
   for v in aa[px:x]:yield v,''
   for v in bb[py:y]:yield '',v
   yield aa[x],bb[y];px=x+1;py=y+1
  for v in aa[px:]:yield v,''
  for v in bb[py:]:yield '',v

def chunks(a,b):
 if not a or not b:return [[('delete',a),('insert',b)]]
 aa=re.findall(r'\S+\s*',a);bb=re.findall(r'\S+\s*',b)
 ops=SequenceMatcher(None,aa,bb,autojunk=False).get_opcodes()
 # Merge tiny unchanged islands into one readable replacement phrase.
 merged=[]
 for z,op in enumerate(ops):
  kind,i,j,k,l=op
  if kind=='equal' and j-i<=2 and z>0 and z+1<len(ops):kind='replace'
  if merged and kind!='equal' and merged[-1][0]!='equal':
   prev=merged.pop();merged.append(('replace',prev[1],j,prev[3],l))
  else:merged.append((kind,i,j,k,l))
 groups=[];cur=[]
 for z,(kind,i,j,k,l) in enumerate(merged):
  if kind=='equal':
   if z==0:
    if j-i>5:cur.append(('equal','… '))
    cur.append(('equal',''.join(aa[max(i,j-5):j])))
   elif z==len(merged)-1:
    cur.append(('equal',''.join(aa[i:min(j,i+5)])))
    if j-i>5:cur.append(('equal',' …'))
   elif j-i>10:
    cur.append(('equal',''.join(aa[i:i+5])+' …'));groups.append(cur)
    cur=[('equal','… '+''.join(aa[j-5:j]))]
   else:cur.append(('equal',''.join(aa[i:j])))
  else:
   if j>i:cur.append(('delete',''.join(aa[i:j])))
   if l>k:cur.append(('insert',''.join(bb[k:l])))
 if cur:groups.append(cur)
 return groups

for section in order:
 pairs=list(aligned(old.get(section,[]),new.get(section,[])))
 if not pairs:continue
 heading=doc.add_paragraph(style='Heading 1');run(heading,section,'insert' if section not in old else 'equal');htmlparts.append('<h2>'+('<ins>'+escape(section)+'</ins>' if section not in old else escape(section))+'</h2>')
 for a,b in pairs:
  ledger.append({'section':section,'removed':a,'added':b})
  for chunk in chunks(a,b):
   p=doc.add_paragraph();p.paragraph_format.widow_control=True;p.paragraph_format.keep_together=True;parts=[]
   for n,(kind,value) in enumerate(chunk):
    if not value:continue
    if n and chunk[n-1][0]=='delete' and kind=='insert':
     p.add_run(' → ');parts.append(' → ')
    run(p,value,kind)
    parts.append(('<del>'+escape(value)+'</del>') if kind=='delete' else ('<ins>'+escape(value)+'</ins>') if kind=='insert' else escape(value))
   htmlparts.append('<p>'+''.join(parts)+'</p>')
doc.add_paragraph('Formát',style='Heading 1')
p=doc.add_paragraph();note='Přibyl editovatelný DOCX s textem, tabulkami a šesti ilustracemi. Týmová sekce posunula stránkování.'
run(p,note,'insert');htmlparts+=['<h2>Formát</h2><p><ins>'+note+'</ins></p>']
footer=page.footer.paragraphs[0];footer.alignment=2
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.save(D/'private-lending-aleo-grant-changes-v3-to-v8.docx')
(D/'private-lending-aleo-grant-changes-v3-to-v8.html').write_text('\n'.join(htmlparts))
(D/'research-2026-10-07/grant-v8-redline.json').write_text(json.dumps(ledger,ensure_ascii=False,indent=2))
print('v8 redline created',len(ledger),'changes')
