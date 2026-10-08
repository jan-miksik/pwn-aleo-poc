"""Keep source prose editable and carry the original illustrations into DOCX."""
from pathlib import Path
from lxml import html
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT

DOCS=Path(__file__).resolve().parent.parent
WIDTH=17.2
INK='17191C'; MUTED='4A4F57'; TEAL='2E6E7E'; BLUE='1F3A5F'

def hexcolor(value):return RGBColor.from_string(value)

def inline(p,node,bold=False,size=10.2,color=INK):
    def text(value,weight=bold):
        if value:
            r=p.add_run(value);r.bold=weight;r.font.name='Arial';r.font.size=Pt(size);r.font.color.rgb=hexcolor(color)
    text(node.text)
    for child in node:
        if child.tag in ['br','wbr']:
            if child.tag=='br':p.add_run().add_break()
            text(child.text)
        elif child.tag=='a':
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(child.get('href'),RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),TEAL);pr.append(c)
            s=OxmlElement('w:sz');s.set(qn('w:val'),str(round(size*2)));pr.append(s)
            r.append(pr);t=OxmlElement('w:t');t.text=''.join(child.itertext());r.append(t);h.append(r);p._p.append(h)
        else:
            inline(p,child,bold or child.tag in ['b','strong'] or 't' in (child.get('class') or '').split(),size,color)
        text(child.tail)

def cell_borders(cell,color='D9D9D9',fill=None):
    pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),color);b.append(x)
    pr.append(b)
    margin=OxmlElement('w:tcMar')
    for side in ['top','bottom','left','right']:
        x=OxmlElement('w:'+side);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');margin.append(x)
    pr.append(margin)
    if fill:
        s=OxmlElement('w:shd');s.set(qn('w:fill'),fill);pr.append(s)

class Builder:
    def __init__(self):
        self.d=Document();self.pending_break=False
        s=self.d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
        s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(1.9);s.right_margin=Cm(1.9)
        s.footer_distance=Cm(.8)
        st=self.d.styles['Normal'];st.font.name='Arial';st.font.size=Pt(10.2)
        st.paragraph_format.line_spacing=1.35;st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.widow_control=True
        for name,size in [('Title',28),('Heading 1',16),('Heading 2',11.5),('Heading 3',9)]:
            st=self.d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=hexcolor(INK)
            st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12 if name!='Title' else 0);st.paragraph_format.space_after=Pt(7)
        p=s.footer.paragraphs[0];p.alignment=2;r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
        self.d.core_properties.title='Private Lending for Aleo'
        self.d.core_properties.author='PWN DAO Foundation'
        self.d.settings.element.append(OxmlElement('w:updateFields'))
        self.d.settings.element[-1].set(qn('w:val'),'true')
    def paragraph(self,node=None,style=None,size=10.2,color=INK,bold=False):
        p=self.d.add_paragraph(style=style)
        if self.pending_break:p.paragraph_format.page_break_before=True;self.pending_break=False
        if node is not None:inline(p,node,bold,size,color)
        return p
    def image(self,name):
        p=self.paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(5)
        shape=p.add_run().add_picture(str(DOCS/'grant-assets-v6'/(name+'.png')),width=Cm(WIDTH))
        shape._inline.docPr.set('descr',name.replace('-',' '))
    def native_table(self,rows,widths=None):
        t=self.d.add_table(rows=len(rows),cols=len(rows[0]));t.autofit=False
        widths=widths or [WIDTH/len(rows[0])]*len(rows[0])
        for j,w in enumerate(widths):t.columns[j].width=Cm(w)
        for i,row in enumerate(rows):
            tr=t.rows[i];tr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if i==0:tr._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for j,node in enumerate(row):
                c=tr.cells[j];c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell_borders(c,fill='F5F6F8' if i==0 else None)
                p=c.paragraphs[0];p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(2)
                inline(p,node,bold=i==0,size=8.6 if len(row)==4 else 9.1)
        self.paragraph().paragraph_format.space_after=Pt(2)
    def heading(self,node):
        level={'h1':'Title','h2':'Heading 1','h3':'Heading 2','h4':'Heading 3'}[node.tag]
        size={'h1':28,'h2':16,'h3':11.5,'h4':9}[node.tag]
        p=self.paragraph(node,level,size,bold=True)
        if node.tag=='h2':
            pp=p._p.get_or_add_pPr();b=OxmlElement('w:pBdr');top=OxmlElement('w:top')
            top.set(qn('w:val'),'single');top.set(qn('w:sz'),'8');top.set(qn('w:color'),INK);top.set(qn('w:space'),'9');b.append(top);pp.append(b)
    def visit(self,node):
        if not isinstance(node.tag,str):return
        cls=set((node.get('class') or '').split())
        if 'page-break' in cls:self.pending_break=True
        if node.tag in ['h1','h2','h3','h4']:self.heading(node);return
        if node.tag=='p':
            self.paragraph(node,size=11 if 'lead' in cls else 9 if 'small' in cls else 10.2,color=MUTED if 'small' in cls or 'muted' in cls else INK);return
        if node.tag=='table':
            rows=[r.xpath('./th|./td') for r in node.xpath('.//tr')]
            self.native_table(rows,[4,6.6,6.6] if len(rows[0])==3 else None);return
        if 'masthead' in cls:
            for c in node:
                if c.tag=='h1':self.heading(c)
                elif 'kicker' in (c.get('class') or ''):self.paragraph(c,size=8.7,color=TEAL,bold=True)
                else:self.paragraph(c,size=11.5,color=MUTED)
            return
        if 'facts' in cls:
            facts=node.xpath('./div');t=self.d.add_table(rows=1,cols=len(facts));t.autofit=False
            for i,f in enumerate(facts):
                c=t.cell(0,i);c.width=Cm(WIDTH/len(facts));cell_borders(c)
                lab=f.xpath('./div[@class="l"]')[0];v=f.xpath('./div[@class="v"]')[0]
                inline(c.paragraphs[0],lab,size=7.8,color=MUTED)
                p=c.add_paragraph();p.paragraph_format.space_after=Pt(3)
                lead=html.Element('b');lead.text=v.text;inline(p,lead,True,12)
                for small in v.xpath('./small'):
                    p=c.add_paragraph();p.paragraph_format.space_after=Pt(1);inline(p,small,size=8.5,color=MUTED)
            self.paragraph();return
        if 'status' in cls:
            boxes=node.xpath('./div');t=self.d.add_table(rows=1,cols=len(boxes));t.autofit=False
            for i,x in enumerate(boxes):
                c=t.cell(0,i);c.width=Cm(WIDTH/len(boxes));cell_borders(c)
                title=x.find('b');p=c.paragraphs[0];p.paragraph_format.line_spacing=1.2;inline(p,title,bold=True,size=9.5)
                p.paragraph_format.space_after=Pt(4)
                p=c.add_paragraph();p.paragraph_format.line_spacing=1.25;p.paragraph_format.space_after=Pt(1)
                for child in x:
                    if child is title:continue
                    inline(p,child,bold=True,size=9,color=TEAL)
                    if child.tail:
                        r=p.add_run(child.tail);r.font.size=Pt(9);r.font.name='Arial'
            self.paragraph();return
        if 'cols' in cls and node.xpath('.//h4[text()="Public by design"]'):
            t=self.d.add_table(rows=1,cols=2);t.autofit=False
            for i,col in enumerate(node.xpath('./div')):
                cell=t.cell(0,i);cell.width=Cm(WIDTH/2);cell_borders(cell)
                for j,child in enumerate(col):
                    p=cell.paragraphs[0] if j==0 else cell.add_paragraph()
                    p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(5)
                    if child.tag=='h4':inline(p,child,bold=True,size=8.8,color=MUTED)
                    elif child.get('class')=='chips':
                        for k,chip in enumerate(child):
                            if k:p.add_run(' · ')
                            inline(p,chip,size=8.8,color=BLUE if i else MUTED)
                    else:inline(p,child,size=9,color=MUTED)
            self.paragraph().paragraph_format.space_after=Pt(0)
            return
        if 'draws' in cls:self.image('elastic-draws');return
        if 'timeline' in cls:self.image('timeline');return
        if node.tag=='svg':
            self.image('architecture' if node.get('viewbox')=='0 0 760 428' else 'lifecycle');return
        if 'flow' in cls:
            self.image('privacy-boundary' if 'Where the boundary' in ''.join(node.getparent().itertext()) else 'discovery');return
        if 'stack' in cls:
            rows=[r.xpath('./div') for r in node.xpath('./div')];self.native_table(rows,[3,4.9,3.6,5.7]);return
        if 'chips' in cls:
            for chip in node:
                p=self.paragraph(chip,size=9)
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.left_indent=Cm(.15)
                pr=p._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6EEF7' if 'private' in (chip.get('class') or '') else 'ECEEF2');pr.append(sh)
            return
        if node.tag in ['ol','ul']:
            for i,li in enumerate(node.xpath('./li'),1):
                p=self.paragraph();p.paragraph_format.left_indent=Cm(.6);p.paragraph_format.first_line_indent=Cm(-.6)
                r=p.add_run(f'{i:02d}  ' if node.tag=='ol' else '•  ');r.font.size=Pt(9)
                inline(p,li,size=9.6)
            return
        if 'ms' in cls:
            title=node.xpath('./div[2]/div[@class="title"]')[0];num=node.xpath('./div[@class="n"]')[0].text
            h=html.Element('h3');h.text=num+'  '+''.join(title.itertext());self.heading(h)
            meta=node.xpath('./div[@class="meta"]')[0]
            label=html.Element('span');label.text=meta.find('b').text+' · '+(meta.find('b').tail or '').strip()
            self.paragraph(label,size=9.5,color=MUTED,bold=True).paragraph_format.keep_with_next=True
            for p in node.xpath('./div[2]/p'):self.visit(p)
            return
        if 'ref' in cls:
            first=node.xpath('./div')[0]
            for c in first:self.paragraph(c,size=9.5,bold=c.get('class')=='n').paragraph_format.keep_with_next=True
            self.paragraph(node.xpath('./div')[1],size=9.5)
            return
        if 'cap' in cls:self.paragraph(node,size=8.5,color=MUTED);return
        if node.tag in ['style','script','defs','marker','path'] or 'space' in cls:return
        for c in node:self.visit(c)

def markdown(root):
    chunks=[]
    def walk(e):
        if not isinstance(e.tag,str):return
        if e.tag=='svg':return
        if e.tag in ['h1','h2','h3','h4']:
            chunks.append('#'*int(e.tag[1])+' '+ ' '.join(e.text_content().split())+'\n');return
        if e.tag=='p':
            def md(x):
                text=x.text or ''
                for c in x:
                    v=md(c)
                    if c.tag in ['strong','b']:v='**'+v+'**'
                    elif c.tag=='a':v='['+v+']('+c.get('href')+')'
                    elif c.tag=='br':v=' '
                    text+=v+(c.tail or '')
                return text
            chunks.append(md(e).strip()+'\n');return
        if e.tag=='table':
            rows=e.xpath('.//tr');lines=[]
            for i,r in enumerate(rows):
                cells=[' '.join(c.text_content().split()) for c in r.xpath('./th|./td')]
                lines.append('| '+' | '.join(cells)+' |')
                if i==0:lines.append('| '+' | '.join(['---']*len(cells))+' |')
            chunks.append('\n'.join(lines)+'\n');return
        if e.tag in ['ul','ol']:
            for i,c in enumerate(e,1):chunks.append(('- ' if e.tag=='ul' else str(i)+'. ')+' '.join(c.text_content().split()))
            chunks.append('');return
        cls=set((e.get('class') or '').split())
        if cls.intersection(['draws','timeline','stack','facts','status','chips','flow']):
            chunks.append(' '.join(e.text_content().split())+'\n');return
        if 'ref' in cls:
            chunks.append(' '.join(e.text_content().split())+'\n');return
        if e.tag not in ['style','script']:
            for c in e:walk(c)
    walk(root.find('body'))
    return '\n'.join(chunks)

if __name__=='__main__':
    root=html.fromstring((DOCS/'private-lending-aleo-grant-overview-v6.html').read_text())
    b=Builder();b.visit(root.find('body'))
    path=DOCS/'private-lending-aleo-grant-draft-v6.docx';b.d.save(path);print(path)
    (DOCS/'private-lending-aleo-grant-draft-v6.md').write_text(markdown(root))
