"""Only changed content blocks, compared directly with the pre-edit v3 source."""
from pathlib import Path
from difflib import SequenceMatcher
from copy import deepcopy
from lxml import html
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

D=Path(__file__).resolve().parent.parent
def blocks(path):
    root=html.fromstring(path.read_text());out=[];section='Úvod'
    def text(e):
        copy=deepcopy(e)
        for child in copy.iterdescendants():
            if child.tag in ['div','small','td','th','br'] or ('axis' in (copy.get('class') or '') and child.tag=='span'):
                child.text=' '+(child.text or '')
                child.tail=' '+(child.tail or '')
        value=' '.join(''.join(copy.itertext()).split())
        links=[]
        for a in e.xpath('.//a[@href]'):
            if a.get('href') not in links:links.append(a.get('href'))
        if links:value+=' ['+'; '.join(links)+']'
        return value
    def visit(e):
        nonlocal section
        if not isinstance(e.tag,str) or e.tag in ['style','script','defs']:return
        cls=set((e.get('class') or '').split())
        if e.tag=='h2':section=text(e)
        if e.tag in ['h1','h2','h3','h4','p','li','tr','text'] or cls.intersection({'kicker','sub','subtitle','fact','chip','ref','cap','box','loan','lender','axis','row','title','meta'}):
            value=text(e)
            if value:out.append((section,value))
            return
        for c in e:visit(c)
    visit(root.find('body'));return out

before=blocks(D/'private-lending-aleo-grant-overview-v3.html')
after=blocks(D/'private-lending-aleo-grant-overview-v6.html')
matcher=SequenceMatcher(None,[x[1] for x in before],[x[1] for x in after],autojunk=False)
changes=[]
for op,i,j,k,l in matcher.get_opcodes():
    if op=='equal':continue
    changes.append({'section':(after[k][0] if k<l else before[i][0]),'removed':[x[1] for x in before[i:j]],'added':[x[1] for x in after[k:l]]})

doc=Document();s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.8);s.left_margin=s.right_margin=Cm(2)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
normal.paragraph_format.line_spacing=1.18;normal.paragraph_format.space_after=Pt(6)
for style,size in [('Title',22),('Heading 1',14),('Heading 2',11)]:
    st=doc.styles[style];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('17191C')
    st.paragraph_format.keep_with_next=True
doc.core_properties.title='Změny grantového návrhu oproti původní verzi'
doc.add_paragraph('Změny grantového návrhu oproti původní verzi',style='Title')
doc.add_paragraph('Původní draft před dnešními úpravami (v3) → aktuální verze (v6). Pouze změněné, přidané a odebrané pasáže; nezměněný obsah je vynechán.')
md=['# Změny grantového návrhu oproti původní verzi','', 'Původní draft před dnešními úpravami (v3) → aktuální verze (v6). Pouze změny.','']
previous=None
for index,change in enumerate(changes,1):
    if previous!=change['section']:
        doc.add_paragraph(change['section'],style='Heading 1');md+=['## '+change['section'],''];previous=change['section']
    for key,label,color in [('removed','Odebráno nebo nahrazeno','9A3434'),('added','Nové znění nebo doplnění','216453')]:
        if not change[key]:continue
        p=doc.add_paragraph(label,style='Heading 2');p.runs[0].font.color.rgb=RGBColor.from_string(color)
        md+=['**'+label+'**','']
        for value in change[key]:
            p=doc.add_paragraph(value);p.paragraph_format.widow_control=True
            md+=[value,'']

# Non-text changes are intentionally listed without repeating the whole design.
doc.add_paragraph('Rozložení a editovatelný formát',style='Heading 1')
layout='Na začátku přibyla týmová sekce, takže se posunulo stránkování. Původní ilustrace zůstaly; popisek cíle časové osy se změnil z „Capped mainnet“ na „Mainnet“. Přibyl DOCX s editovatelným textem a tabulkami a se šesti vloženými ilustracemi.'
doc.add_paragraph(layout);md+=['## Rozložení a editovatelný formát','',layout,'']
footer=s.footer.paragraphs[0];footer.alignment=2
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.save(D/'private-lending-aleo-grant-changes-v3-to-v6.docx')
(D/'private-lending-aleo-grant-changes-v3-to-v6.md').write_text('\n'.join(md))
(D/'research-2026-10-07/grant-v6-content-changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2))
print(json.dumps({'changed_blocks':len(changes),'removed':sum(len(x['removed']) for x in changes),'added':sum(len(x['added']) for x in changes)}))
